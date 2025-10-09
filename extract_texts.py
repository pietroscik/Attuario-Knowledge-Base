#!/usr/bin/env python3
# ==========================================================
#  EXTRACT_TEXTS.PY
#  Estrae testo da materiali PDF, PPTX e script R
#  per la Attuario Knowledge Base.
#  Autore: pietroscik · 2025
# ==========================================================

import os
import re
import csv
from pathlib import Path
from datetime import datetime

import fitz  # PyMuPDF (per PDF)
from pptx import Presentation  # per PowerPoint

# OCR opzionale (per PDF scansionati)
try:
    import pytesseract
    from PIL import Image
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False

# === Percorsi base ============================================================
BASE_DIR = Path(__file__).resolve().parent
RAW_DIR = BASE_DIR / "01_raw_materials"
OUTPUT_DIR = BASE_DIR / "02_text_extracted" / "parsed_texts"
INDEX_FILE = BASE_DIR / "materials_index.csv"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# ==============================================================================

def extract_text_from_pdf(file_path: Path) -> str:
    """Estrae testo da file PDF, con fallback OCR se necessario."""
    text = ""
    try:
        with fitz.open(file_path) as doc:
            for page in doc:
                text += page.get_text("text")
        if not text.strip() and OCR_AVAILABLE:
            for page_number in range(len(doc)):
                pix = doc.load_page(page_number).get_pixmap()
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                text += pytesseract.image_to_string(img)
    except Exception as e:
        text = f"[ERRORE PDF: {e}]"
    return text.strip()

# ==============================================================================

def extract_text_from_pptx(file_path: Path) -> str:
    """Estrae testo da una presentazione PowerPoint (.pptx)."""
    text_runs = []
    try:
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
    except Exception as e:
        return f"[ERRORE PPTX: {e}]"
    return "\n".join(text_runs).strip()

# ==============================================================================

def extract_text_from_script(file_path: Path) -> str:
    """Legge semplicemente il contenuto di file R o Python."""
    try:
        return Path(file_path).read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        return f"[ERRORE SCRIPT: {e}]"

# ==============================================================================

def sanitize_filename(name: str) -> str:
    """Rimuove caratteri non validi per nomi file."""
    return re.sub(r"[^a-zA-Z0-9_-]", "_", name)

# ==============================================================================

def update_index(material_type: str, src_path: Path, out_path: Path, text_len: int):
    """Aggiorna (o crea) il file CSV materials_index.csv."""
    header = ["timestamp", "material_type", "source_file", "output_file", "char_count"]
    new_row = [datetime.now().isoformat(), material_type, str(src_path), str(out_path), text_len]

    file_exists = INDEX_FILE.exists()
    with open(INDEX_FILE, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if not file_exists:
            writer.writerow(header)
        writer.writerow(new_row)

# ==============================================================================

def main():
    print("📘 Estrazione testi in corso...")

    material_dirs = {
        "pdf": RAW_DIR / "lezioni_pdf",
        "pptx": RAW_DIR / "slide_ppt",
        "r": RAW_DIR / "script_r",
    }

    for mtype, directory in material_dirs.items():
        for file in directory.glob(f"*.{mtype}"):
            print(f" → {file.name}")
            if mtype == "pdf":
                text = extract_text_from_pdf(file)
            elif mtype == "pptx":
                text = extract_text_from_pptx(file)
            elif mtype == "r":
                text = extract_text_from_script(file)
            else:
                continue

            output_name = sanitize_filename(file.stem) + ".txt"
            output_path = OUTPUT_DIR / output_name

            output_path.write_text(text, encoding="utf-8")
            update_index(mtype, file, output_path, len(text))

    print("✅ Estrazione completata. File salvati in 02_text_extracted/parsed_texts/")

# ==============================================================================

if __name__ == "__main__":
    main()
